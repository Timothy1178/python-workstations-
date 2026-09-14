"""Turn chat-box attachments into prompt context and image files.

The shared chat box (static/chatbox.js) sends attachments as
[{name, text}] for text files read in the browser, or [{name, data}]
with base64 for binary files. This module extracts what each provider
can use:
- text files and Excel workbooks become inline document blocks in the
  prompt (openpyxl renders sheets as CSV-style text);
- images are saved under data/uploads/ and their paths handed to the
  provider — the local Claude CLI reads them with its Read tool;
  Copilot Studio agents can't receive local files, so the prompt says so;
- unsupported binaries (pdf, docx, …) get a note asking for text.
"""
import base64
import io
import re
import time
import uuid
from pathlib import Path

UPLOAD_DIR = Path(__file__).resolve().parents[2] / "data" / "uploads"

TEXT_EXT = {"txt", "md", "csv", "tsv", "json", "log", "xml", "yml", "yaml",
            "html", "htm", "sql", "py", "js", "ini", "cfg", "properties"}
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}
EXCEL_EXT = {"xlsx", "xlsm", "xltx"}

MAX_FILE_CHARS = 25000     # per attached document
MAX_TOTAL_CHARS = 60000    # all documents together
MAX_SHEET_ROWS = 300


def _ext(name):
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _safe_name(name):
    base = re.sub(r"[^\w.\-]+", "_", name)[-80:] or "file"
    return f"{time.strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:6]}-{base}"


def _clip(text, note_name):
    if len(text) > MAX_FILE_CHARS:
        return (text[:MAX_FILE_CHARS]
                + f"\n[... {note_name} truncated at {MAX_FILE_CHARS} characters]")
    return text


def _xlsx_to_text(raw):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        lines = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                lines.append(f"[... more rows truncated at {MAX_SHEET_ROWS}]")
                break
            lines.append(",".join("" if c is None else str(c) for c in row))
        parts.append(f"[sheet: {ws.title}]\n" + "\n".join(lines))
    wb.close()
    return "\n\n".join(parts)


def process(attachment_list):
    """Returns {"context": str, "image_paths": [str, ...]}."""
    blocks, image_paths, total = [], [], 0
    for att in attachment_list or []:
        name = str(att.get("name") or "file")
        ext = _ext(name)
        text = att.get("text")
        raw = None
        if text is None and att.get("data"):
            try:
                raw = base64.b64decode(att["data"])
            except Exception:
                blocks.append(f"--- {name} ---\n[could not decode this file]")
                continue

        if text is not None:
            body = _clip(str(text), name)
        elif raw is not None and ext in IMAGE_EXT:
            UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
            path = UPLOAD_DIR / _safe_name(name)
            path.write_bytes(raw)
            image_paths.append(str(path))
            continue
        elif raw is not None and ext in EXCEL_EXT:
            try:
                body = _clip(_xlsx_to_text(raw), name)
            except Exception as exc:
                body = f"[could not read this Excel file: {exc}]"
        elif raw is not None and ext in TEXT_EXT:
            body = _clip(raw.decode("utf-8-sig", errors="replace"), name)
        else:
            body = ("[this file type is not supported yet — paste its text "
                    "into the chat, or convert it to txt/csv/xlsx]")

        total += len(body)
        if total > MAX_TOTAL_CHARS:
            blocks.append(f"--- {name} ---\n[skipped — attached documents "
                          f"exceed {MAX_TOTAL_CHARS} characters in total]")
            continue
        blocks.append(f"--- {name} ---\n{body}")

    context = ""
    if blocks:
        context = ("The tester attached these documents — use them as "
                   "source material:\n\n" + "\n\n".join(blocks) + "\n\n")
    return {"context": context, "image_paths": image_paths}
