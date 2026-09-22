"""Export a test plan to Excel in the CCASIA/AIMAS test plan format.

Starts from assets/base_template.xlsx — a cleaned copy of the user's own
"CCASIA-44 UAT Test Plan.xlsx" — so the summary formulas (COUNTA/COUNTIF),
header styling, merged description cell, column widths and freeze panes are
the original ones, untouched. This module only fills in data rows and
rebuilds the ScreenCap sheet.

Screenshots are embedded at native pixel data (openpyxl never recompresses;
display size is capped to fit the sheet but the stored image keeps full
resolution, so zooming in stays sharp).
"""
import io
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image as PILImage

from . import store

BASE_TEMPLATE = Path(__file__).parent / "assets" / "base_template.xlsx"

# SIT sheet columns A..J -> where each value comes from
COLUMNS = [
    ("testing_bu", "field"),
    ("case_id", "field"),
    ("title", "field"),
    ("steps", "field"),
    ("expected", "field"),
    ("result", "run"),
    ("actual", "run"),
    ("test_data", "field"),
    ("campaign_id", "field"),
    ("remark", "field"),
]
WRAP_COLS = {3, 4, 5, 7, 8, 10}  # C, D, E, G, H, J

FIRST_DATA_ROW = 9

# ScreenCap layout: labels in column A, images anchored at column B.
# Excel's default row height ≈ 20 px; used to advance past each image.
PX_PER_ROW = 20.0
MAX_IMG_WIDTH = 1000  # display cap in px — pixel data is kept as-is
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
BAND_FILL = PatternFill("solid", fgColor="DDEBF7")


def export_plan(plan):
    wb = openpyxl.load_workbook(BASE_TEMPLATE)
    sit = wb["SIT"]

    sit["B1"] = plan.get("description") or plan.get("name", "")

    anchors = _build_screencap(wb["ScreenCap"], plan)

    body_font = Font(name="Calibri", size=11)
    for i, case in enumerate(plan.get("cases", [])):
        row = FIRST_DATA_ROW + i
        for col, (key, source) in enumerate(COLUMNS, start=1):
            if source == "run":
                value = case.get("run", {}).get(key, "")
                if key == "result" and value == "Not Run":
                    value = ""
            else:
                value = case.get("fields", {}).get(key, "")
            cell = sit.cell(row=row, column=col, value=value or None)
            cell.font = body_font
            cell.alignment = Alignment(
                vertical="top", wrap_text=(col in WRAP_COLS))
        # Case ID links to the case's evidence block on the ScreenCap sheet.
        anchor = anchors.get(case.get("uid"))
        if anchor:
            id_cell = sit.cell(row=row, column=2)
            if id_cell.value is None:
                id_cell.value = str(i + 1)
            id_cell.hyperlink = Hyperlink(ref=id_cell.coordinate,
                                          location=f"'ScreenCap'!A{anchor}")
            id_cell.font = LINK_FONT
            anchors[case["uid"]] = (anchor, row)
    for uid, (shot_row, sit_row) in list(anchors.items()):
        back = wb["ScreenCap"].cell(row=shot_row, column=3, value="↑ back to SIT")
        back.hyperlink = Hyperlink(ref=back.coordinate, location=f"'SIT'!B{sit_row}")
        back.font = LINK_FONT

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_screencap(ws, plan):
    """One band per case (label row, then numbered captions + images);
    returns {case uid: band row} so the SIT sheet can link to it."""
    label_font = Font(name="Calibri", size=12, bold=True)
    caption_font = Font(name="Calibri", size=10, italic=True, color="595959")
    anchors = {}
    row = 1
    for i, case in enumerate(plan.get("cases", []), start=1):
        shots = store.list_shots(plan["id"], case["uid"])
        if not shots:
            continue
        bu = case.get("fields", {}).get("testing_bu", "").strip()
        title = case.get("fields", {}).get("title", "").strip()
        num = case.get("fields", {}).get("case_id", "").strip() or str(i)
        label = " ".join(x for x in (bu, title) if x) + f" #{num}"
        anchors[case["uid"]] = row
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).fill = BAND_FILL
        cell = ws.cell(row=row, column=1, value=f"📷 {label.strip()}")
        cell.font = label_font
        cell.alignment = Alignment(vertical="center")
        ws.cell(row=row, column=2,
                value=f"{len(shots)} evidence image{'s' if len(shots) > 1 else ''}"
                ).font = caption_font
        row += 1

        for n, name in enumerate(shots, start=1):
            cap = ws.cell(row=row, column=2, value=f"Evidence {n} of {len(shots)} — {name}")
            cap.font = caption_font
            row += 1
            path = store.shots_dir(plan["id"], case["uid"]) / name
            img = XLImage(str(path))
            with PILImage.open(path) as pil:
                w, h = pil.size
            # Cap the DISPLAY size only; the embedded pixels stay full-res.
            if w > MAX_IMG_WIDTH:
                scale = MAX_IMG_WIDTH / w
                img.width, img.height = MAX_IMG_WIDTH, h * scale
            else:
                img.width, img.height = w, h
            img.anchor = f"B{row}"
            ws.add_image(img)
            row += int(img.height / PX_PER_ROW) + 2
        row += 1  # gap between cases
    return anchors
