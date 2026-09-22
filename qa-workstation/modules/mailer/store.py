"""Storage for the Mailer: templates, receiver groups, jobs, settings, log.

Everything lives under data/mailer/ as JSON. A job's data source is one
of: a saved table (captured from the clipboard by the plugin, editable),
a local .xlsx path (e.g. a OneDrive-synced online workbook) or a direct
download URL — the last two are re-read at send time, so scheduled mails
carry live data.
"""
import html
import io
import json
import re
import time
import uuid
from pathlib import Path

DATA_DIR = Path(__file__).resolve().parents[2] / "data" / "mailer"
FILES = {"templates": "templates.json", "receivers": "receivers.json",
         "jobs": "jobs.json", "settings": "settings.json", "log": "log.json"}

DEFAULT_TEMPLATE_HTML = """<p>Hi all,</p>
<p>Please find the latest <b>{{job}}</b> as of {{date}}:</p>
{{table}}
<p style="color:#6b7280;font-size:12px">{{rows}} rows · sent automatically by the QA Workstation Mailer.</p>"""


def _path(kind):
    return DATA_DIR / FILES[kind]


def _load(kind, default):
    p = _path(kind)
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return default


def _save(kind, data):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    _path(kind).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def _new_id():
    return f"{time.strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:5]}"


def now():
    return time.strftime("%Y-%m-%d %H:%M")


# ---------- generic list helpers ----------

def list_items(kind):
    return _load(kind, [])


def get_item(kind, item_id):
    return next((x for x in list_items(kind) if x["id"] == item_id), None)


def upsert(kind, item):
    items = list_items(kind)
    if not item.get("id"):
        item["id"] = _new_id()
        item["created"] = now()
        items.insert(0, item)
    else:
        for i, x in enumerate(items):
            if x["id"] == item["id"]:
                items[i] = {**x, **item}
                break
        else:
            items.insert(0, item)
    item["updated"] = now()
    _save(kind, items)
    return item


def delete_item(kind, item_id):
    _save(kind, [x for x in list_items(kind) if x["id"] != item_id])


# ---------- settings & log ----------

def settings():
    return {"transport": "outlook", "smtp_host": "smtp.office365.com", "smtp_port": 587,
            "smtp_user": "", "smtp_password": "", "from_name": "", "review_before_send": False,
            **_load("settings", {})}


def save_settings(data):
    cur = settings()
    cur.update({k: v for k, v in data.items() if k in cur})
    _save("settings", cur)
    return cur


def log(entry):
    entries = _load("log", [])
    entries.insert(0, {"at": now(), **entry})
    _save("log", entries[:200])


def recent_log(n=30):
    return _load("log", [])[:n]


# ---------- tables ----------

def parse_tsv(text):
    rows = [r.split("\t") for r in text.replace("\r", "").split("\n") if r.strip()]
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def parse_email_list(raw):
    return [e.strip() for e in re.split(r"[,;\s]+", raw or "") if e.strip()]


def table_html(table, style=True):
    """{"headers": [...], "rows": [[...]]} → an email-safe HTML table with
    inline styles (Outlook ignores stylesheets)."""
    if not table:
        return ""
    headers = table.get("headers") or []
    rows = table.get("rows") or []
    th = ("border:1px solid #cbd5e1;padding:6px 10px;background:#e0ecff;"
          "font-weight:600;text-align:left;font:13px Calibri,Arial,sans-serif;")
    td = "border:1px solid #cbd5e1;padding:5px 10px;font:13px Calibri,Arial,sans-serif;"
    out = ['<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;border:1px solid #cbd5e1;">']
    if headers:
        out.append("<tr>" + "".join(f'<th style="{th if style else ""}">{html.escape(str(h))}</th>' for h in headers) + "</tr>")
    for i, r in enumerate(rows):
        bg = "background:#f8fafc;" if i % 2 else ""
        out.append("<tr>" + "".join(f'<td style="{(td + bg) if style else ""}">{html.escape(str(c))}</td>' for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def read_xlsx_range(raw_bytes, sheet="", cell_range="", header_row=True):
    """Extract a sheet/range from workbook bytes → {"headers", "rows"}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    rows = []
    if cell_range:
        cells = ws[cell_range]
        if not isinstance(cells, tuple):
            cells = ((cells,),)
        elif cells and not isinstance(cells[0], tuple):
            cells = (cells,)
        for r in cells:
            rows.append(["" if c.value is None else _fmt(c.value) for c in r])
    else:
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else _fmt(v) for v in r])
    rows = [r for r in rows if any(str(v).strip() for v in r)]
    wb.close()
    if not rows:
        return {"headers": [], "rows": []}
    if header_row:
        return {"headers": rows[0], "rows": rows[1:]}
    return {"headers": [], "rows": rows}


def _fmt(v):
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d") if getattr(v, "hour", 0) == 0 and getattr(v, "minute", 0) == 0 else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def resolve_table(job):
    """Fetch the job's data: stored table, local xlsx, or URL."""
    src = job.get("source") or {}
    kind = src.get("type", "table")
    if kind == "file":
        p = Path(src.get("path", "")).expanduser()
        if not p.exists():
            raise FileNotFoundError(f"workbook not found: {p}")
        return read_xlsx_range(p.read_bytes(), src.get("sheet", ""), src.get("range", ""),
                               src.get("header_row", True))
    if kind == "url":
        import requests
        r = requests.get(src.get("url", ""), timeout=60)
        r.raise_for_status()
        return read_xlsx_range(r.content, src.get("sheet", ""), src.get("range", ""),
                               src.get("header_row", True))
    return src.get("table") or {"headers": [], "rows": []}


def render(template, job, table):
    """Fill the template placeholders; returns (subject, html_body)."""
    values = {
        "table": table_html(table),
        "date": time.strftime("%Y-%m-%d"),
        "datetime": now(),
        "job": job.get("name", ""),
        "rows": str(len(table.get("rows") or [])),
    }
    def sub(text):
        return re.sub(r"\{\{\s*(\w+)\s*\}\}", lambda m: values.get(m.group(1), m.group(0)), text or "")
    subject = sub(job.get("subject") or template.get("subject") or job.get("name", "Report"))
    body = sub(template.get("html") or DEFAULT_TEMPLATE_HTML)
    if "{{table}}" not in (template.get("html") or DEFAULT_TEMPLATE_HTML):
        body += "\n" + values["table"]
    return subject, body


def recipients_for(job):
    group = get_item("receivers", job.get("receivers_id", "")) or {}
    to = parse_email_list(job.get("to", "")) or list(group.get("to", []))
    cc = parse_email_list(job.get("cc", "")) or list(group.get("cc", []))
    bcc = parse_email_list(job.get("bcc", "")) or list(group.get("bcc", []))
    return to, cc, bcc
